#!/usr/bin/env python3
"""
KAGARIHI 会計Bot — Excel エクスポートスクリプト
使い方（VM上で実行）:
  sudo -u botuser /opt/KAGARIHI_bookkeeping/venv/bin/python \
    /opt/KAGARIHI_bookkeeping/deploy/export_excel.py

出力先: /opt/KAGARIHI_bookkeeping/export/会計データ_YYYYMMDD.xlsx
"""

import os
import sqlite3
import sys
from datetime import datetime
from pathlib import Path

try:
    import openpyxl
    from openpyxl.styles import (
        Alignment, Border, Font, PatternFill, Side
    )
    from openpyxl.utils import get_column_letter
except ImportError:
    print("openpyxl が見つかりません。以下を実行してください:")
    print("  pip install openpyxl")
    sys.exit(1)

# ── 設定 ──────────────────────────────────────────────────────────────────────
INSTALL_DIR = Path(__file__).parent.parent
DB_PATH     = os.getenv("DB_PATH", str(INSTALL_DIR / "bookkeeping.db"))
EXPORT_DIR  = INSTALL_DIR / "export"

# ── スタイル定数 ──────────────────────────────────────────────────────────────
COLOR_HEADER_DARK   = "1F3864"  # 濃紺（シートヘッダー）
COLOR_HEADER_LIGHT  = "D6E4F0"  # 薄青（サブヘッダー）
COLOR_INCOME        = "E8F5E9"  # 薄緑（収益行）
COLOR_EXPENSE       = "FFF3E0"  # 薄橙（費用行）
COLOR_SETTLED       = "F5F5F5"  # 薄灰（精算済み）
COLOR_BORDER        = "BDBDBD"

FONT_HEADER = Font(name="Meiryo UI", bold=True, color="FFFFFF", size=10)
FONT_TITLE  = Font(name="Meiryo UI", bold=True, size=9)
FONT_BODY   = Font(name="Meiryo UI", size=9)

ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
ALIGN_LEFT   = Alignment(horizontal="left",   vertical="center")
ALIGN_RIGHT  = Alignment(horizontal="right",  vertical="center", wrap_text=False)

def thin_border():
    s = Side(style="thin", color=COLOR_BORDER)
    return Border(left=s, right=s, top=s, bottom=s)

def header_fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color)

def set_col_width(ws, col: int, width: float):
    ws.column_dimensions[get_column_letter(col)].width = width

def write_header_row(ws, row: int, headers: list[tuple[str, float]]):
    """headers: [(label, width), ...]"""
    for col, (label, width) in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col, value=label)
        cell.font   = FONT_HEADER
        cell.fill   = header_fill(COLOR_HEADER_DARK)
        cell.alignment = ALIGN_CENTER
        cell.border = thin_border()
        set_col_width(ws, col, width)

def style_data_row(ws, row: int, n_cols: int, fill_color: str | None = None):
    for col in range(1, n_cols + 1):
        cell = ws.cell(row=row, column=col)
        cell.font   = FONT_BODY
        cell.border = thin_border()
        if fill_color:
            cell.fill = header_fill(fill_color)

# ── シート作成関数 ──────────────────────────────────────────────────────────────

def sheet_仕訳帳(wb, cur):
    ws = wb.create_sheet("仕訳帳")
    ws.freeze_panes = "A2"

    headers = [
        ("ID",        5),
        ("日付",      12),
        ("借方科目",  16),
        ("貸方科目",  16),
        ("金額（税込）", 14),
        ("消費税率",  10),
        ("摘要",      40),
        ("イベント",  16),
        ("登録日時",  18),
    ]
    write_header_row(ws, 1, headers)

    cur.execute("""
        SELECT id, entry_date, debit_account, credit_account,
               amount, tax_rate, description, event_tag, created_at
        FROM journal_entries
        ORDER BY entry_date, id
    """)
    rows = cur.fetchall()

    # 勘定科目→種別のマップ（収益/費用の色分け用）
    cur.execute("SELECT name, account_type FROM accounts")
    acct_type = {r[0]: r[1] for r in cur.fetchall()}

    for r, row in enumerate(rows, 2):
        id_, date_, debit, credit, amount, tax_rate, desc, event, created = row

        debit_type  = acct_type.get(debit, "")
        credit_type = acct_type.get(credit, "")
        if "収益" in (debit_type, credit_type):
            fill = COLOR_INCOME
        elif "費用" in (debit_type, credit_type):
            fill = COLOR_EXPENSE
        else:
            fill = None

        values = [id_, date_, debit, credit, amount,
                  f"{tax_rate}%" if tax_rate else "0%",
                  desc, event or "", created]

        for col, val in enumerate(values, 1):
            cell = ws.cell(row=r, column=col, value=val)
            cell.font   = FONT_BODY
            cell.border = thin_border()
            if fill:
                cell.fill = header_fill(fill)
            if col == 5:  # 金額
                cell.number_format = '#,##0'
                cell.alignment = ALIGN_RIGHT
            elif col in (1, 6):  # ID, 税率
                cell.alignment = ALIGN_CENTER
            else:
                cell.alignment = ALIGN_LEFT

    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"
    return len(rows)


def sheet_立替精算(wb, cur):
    ws = wb.create_sheet("立替精算")
    ws.freeze_panes = "A2"

    headers = [
        ("ID",       5),
        ("日付",    12),
        ("立替者",  14),
        ("金額",    12),
        ("内容",    40),
        ("精算状況", 10),
        ("登録日時", 18),
    ]
    write_header_row(ws, 1, headers)

    cur.execute("""
        SELECT id, entry_date, paid_by, amount, description, settled, created_at
        FROM advances
        ORDER BY entry_date, id
    """)
    for r, row in enumerate(cur.fetchall(), 2):
        id_, date_, paid_by, amount, desc, settled, created = row
        status = "精算済" if settled else "未精算"
        fill   = COLOR_SETTLED if settled else None

        values = [id_, date_, paid_by, amount, desc, status, created]
        for col, val in enumerate(values, 1):
            cell = ws.cell(row=r, column=col, value=val)
            cell.font   = FONT_BODY
            cell.border = thin_border()
            if fill:
                cell.fill = header_fill(fill)
            if col == 4:
                cell.number_format = '#,##0'
                cell.alignment = ALIGN_RIGHT
            elif col in (1, 6):
                cell.alignment = ALIGN_CENTER
            else:
                cell.alignment = ALIGN_LEFT


def sheet_グッズ取引(wb, cur):
    ws = wb.create_sheet("グッズ取引")
    ws.freeze_panes = "A2"

    headers = [
        ("ID",       5),
        ("日付",    12),
        ("グッズ名", 20),
        ("種別",     8),
        ("数量",     8),
        ("単価",    12),
        ("合計金額", 14),
        ("摘要",    30),
        ("登録日時", 18),
    ]
    write_header_row(ws, 1, headers)

    cur.execute("""
        SELECT id, entry_date, goods_name, tx_type, quantity,
               unit_price, total_amount, description, created_at
        FROM goods_transactions
        ORDER BY entry_date, id
    """)
    for r, row in enumerate(cur.fetchall(), 2):
        id_, date_, name, tx_type, qty, unit, total, desc, created = row
        fill = COLOR_INCOME if tx_type == "販売" else COLOR_EXPENSE

        values = [id_, date_, name, tx_type, qty, unit, total, desc, created]
        for col, val in enumerate(values, 1):
            cell = ws.cell(row=r, column=col, value=val)
            cell.font   = FONT_BODY
            cell.fill   = header_fill(fill)
            cell.border = thin_border()
            if col in (6, 7):
                cell.number_format = '#,##0'
                cell.alignment = ALIGN_RIGHT
            elif col in (1, 4, 5):
                cell.alignment = ALIGN_CENTER
            else:
                cell.alignment = ALIGN_LEFT


def sheet_グッズ在庫(wb, cur):
    ws = wb.create_sheet("グッズ在庫")
    ws.freeze_panes = "A2"

    headers = [
        ("グッズ名", 24),
        ("販売単価",  12),
        ("現在庫数",  10),
        ("在庫評価額", 14),
        ("登録日時",  18),
    ]
    write_header_row(ws, 1, headers)

    cur.execute("""
        SELECT name, selling_price, stock, created_at
        FROM goods
        ORDER BY name
    """)
    for r, row in enumerate(cur.fetchall(), 2):
        name, price, stock, created = row
        inventory_val = price * stock

        values = [name, price, stock, inventory_val, created]
        for col, val in enumerate(values, 1):
            cell = ws.cell(row=r, column=col, value=val)
            cell.font   = FONT_BODY
            cell.border = thin_border()
            if col in (2, 4):
                cell.number_format = '#,##0'
                cell.alignment = ALIGN_RIGHT
            elif col == 3:
                cell.alignment = ALIGN_CENTER
            else:
                cell.alignment = ALIGN_LEFT


def sheet_勘定科目(wb, cur):
    ws = wb.create_sheet("勘定科目")
    ws.freeze_panes = "A2"

    headers = [
        ("ID",       5),
        ("勘定科目名", 20),
        ("種別",     10),
    ]
    write_header_row(ws, 1, headers)

    cur.execute("SELECT id, name, account_type FROM accounts ORDER BY account_type, id")
    type_colors = {
        "資産": "E3F2FD", "負債": "FCE4EC",
        "資本": "F3E5F5", "収益": "E8F5E9", "費用": "FFF3E0",
    }
    for r, (id_, name, atype) in enumerate(cur.fetchall(), 2):
        fill = type_colors.get(atype, None)
        for col, val in enumerate([id_, name, atype], 1):
            cell = ws.cell(row=r, column=col, value=val)
            cell.font   = FONT_BODY
            cell.border = thin_border()
            if fill:
                cell.fill = header_fill(fill)
            cell.alignment = ALIGN_CENTER if col in (1, 3) else ALIGN_LEFT


# ── メイン ────────────────────────────────────────────────────────────────────

def main():
    if not Path(DB_PATH).exists():
        print(f"データベースが見つかりません: {DB_PATH}")
        sys.exit(1)

    EXPORT_DIR.mkdir(parents=True, exist_ok=True)

    con = sqlite3.connect(DB_PATH)
    cur = con.cursor()

    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # デフォルトのシートを削除

    # 各シートを作成
    entry_count = sheet_仕訳帳(wb, cur)
    sheet_立替精算(wb, cur)
    sheet_グッズ取引(wb, cur)
    sheet_グッズ在庫(wb, cur)
    sheet_勘定科目(wb, cur)

    con.close()

    today = datetime.now().strftime("%Y%m%d")
    out_path = EXPORT_DIR / f"会計データ_{today}.xlsx"
    wb.save(str(out_path))

    print(f"✅ エクスポート完了: {out_path}")
    print(f"   仕訳帳: {entry_count} 件")


if __name__ == "__main__":
    main()
